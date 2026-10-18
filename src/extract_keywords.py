"""Extract lesson titles and KEY WORDS cards from a unit's lesson decks
(e.g. the Foundation F2 PowerPoints) into src/data/<unit>-keywords.json.
The key words become each Speaking Mastery page's "Key language" station.
Usage: python3 src/extract_keywords.py <folder of decks> <unit code, e.g. F2>
(in the Miftah_arabic repository the F2 folder is Foundation_F2)."""
import glob, json, os, re, sys
from pptx import Presentation

E = 914400
HERE = os.path.dirname(os.path.abspath(__file__))
AR = re.compile(r'[؀-ۿ]')


def shapes(slide):
    return [dict(x=sh.left / E, y=sh.top / E, w=sh.width / E, t=sh.text_frame.text.strip())
            for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def at(S, x, y, tol=0.06):
    return next((s['t'] for s in S if abs(s['x'] - x) < tol and abs(s['y'] - y) < tol), '')


def cards(S):
    items = []
    for a in S:
        if not (abs(a['w'] - 0.34) < 0.03 and a['y'] > 2 and a['t'].isdigit()):
            continue
        x, y = a['x'], a['y']
        head = at(S, x, y + 0.32)
        if not head:
            continue
        below = sorted([s for s in S if abs(s['x'] - x - 0.03) < 0.05 and y + 0.8 < s['y'] < y + 1.6 and not AR.search(s['t'].split('\n')[0])],
                       key=lambda s: s['y'])
        meta = below[0]['t'].split('\n') if below else ['']
        tag = next((s['t'] for s in S if abs(s['y'] - y - 0.02) < 0.04 and x + 0.3 < s['x'] < x + 2.3 and s['t'] not in ('CORE', 'DEVELOP', 'STRETCH')), '')
        forms = []
        for lab in S:
            if abs(lab['y'] - y - 1.67) < 0.05 and x - 0.01 < lab['x'] < x + 3.0 and lab['w'] > 0.9 and not AR.search(lab['t']):
                form = next((s['t'] for s in S if AR.search(s['t']) and abs(s['x'] - lab['x'] + 0.01) < 0.04 and 0 < s['y'] - lab['y'] < 0.3), '')
                if form:
                    forms.append({'label': lab['t'], 'ar': form, 'x': lab['x']})
        forms.sort(key=lambda f: -f['x'])  # cards read right to left: m. · f. · pl.
        for f in forms:
            del f['x']
        items.append({'n': int(a['t']), 'ar': head, 'en': meta[0], 'tr': meta[1] if len(meta) > 1 and not forms else '',
                      'kind': tag, 'forms': forms})
    return sorted(items, key=lambda i: i['n'])


def table_cards(S):
    """Masculine / Feminine / Plural / English tables (e.g. F2-L05)."""
    heads = sorted([h for h in S if abs(h['y'] - 1.98) < 0.05], key=lambda h: h['x'])
    names = [h['t'] for h in heads]
    if 'Masculine' not in names or 'Feminine' not in names:
        return []
    col = lambda s: [h for h in heads if h['x'] <= s['x'] + 0.12][-1]['t'] if any(h['x'] <= s['x'] + 0.12 for h in heads) else None
    rows = {}
    for s in S:
        if s['y'] > 2.3 and s['y'] < 7 and s['t'] not in ('CORE', 'DEVELOP', 'STRETCH'):
            key = round(s['y'] * 4) / 4
            rows.setdefault(min(rows, key=lambda k: abs(k - s['y']), default=key) if rows and min(abs(k - s['y']) for k in rows) < 0.1 else key, {})[col(s)] = s['t']
    items = []
    for i, (_, r) in enumerate(sorted(rows.items())):
        m, f = r.get('Masculine', ''), r.get('Feminine', '')
        if not m:
            continue
        forms = [{'label': 'm.', 'ar': m}, {'label': 'f.', 'ar': f}]
        pl = next((v for k, v in r.items() if k and k.startswith('Plural')), '')
        if pl:
            forms.append({'label': 'pl.', 'ar': pl})
        ex = next((v for k, v in r.items() if k and k.startswith('Website example')), '')
        items.append({'n': i + 1, 'ar': m, 'en': r.get('English', ''), 'tr': '', 'kind': ex, 'forms': forms})
    return items


def lesson(path, unit):
    p = Presentation(path)
    code = re.search(r'(%s-L\d\d)' % unit, os.path.basename(path)).group(1)
    cover = shapes(p.slides[0])
    texts = [s['t'] for s in sorted(cover, key=lambda s: (s['y'], s['x']))]
    i = texts.index(code) if code in texts else 0
    title, arabic, summary = (texts[i + 1:i + 4] + ['', '', ''])[:3]
    bank = []
    for slide in p.slides:
        S = shapes(slide)
        head = at(S, 0.5, 0.9)
        if head.startswith('KEY WORDS'):
            items = cards(S) or table_cards(S)
            if items:
                bank.append({'title': at(S, 0.5, 1.13), 'arabic': at(S, 8.2, 0.95, 0.08), 'items': items})
    return {'code': code, 'title': title, 'arabic': arabic, 'summary': summary, 'bank': bank}


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    folder, unit = sys.argv[1], sys.argv[2].upper()
    files = sorted(glob.glob(os.path.join(folder, '%s-L*.pptx' % unit)))
    data = [lesson(f, unit) for f in files]
    out = os.path.join(HERE, 'data', '%s-keywords.json' % unit.lower())
    json.dump({'unit': unit, 'lessons': data}, open(out, 'w'), ensure_ascii=False, indent=1)
    for L in data:
        print(L['code'], L['title'], '|', L['arabic'], '|', [len(g['items']) for g in L['bank']])
